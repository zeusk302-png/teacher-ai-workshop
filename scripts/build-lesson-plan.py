from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'artifacts/python-libs'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import qrcode

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/docx'; OUT.mkdir(parents=True,exist_ok=True)
URL='https://zeusk302-png.github.io/teacher-ai-workshop/'
qr=ROOT/'docs/assets/course-qr.png'
qrcode.make(URL,border=4,box_size=10).save(qr)
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for key in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=d.styles[key];s.font.name='맑은 고딕';s.font.color.rgb=RGBColor(0,0,0)
 s.element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(4)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(18)
d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.space_after=Pt(4)
d.styles['Subtitle'].font.size=Pt(10)
d.styles['Subtitle'].font.italic=False
d.styles['Subtitle'].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].font.size=Pt(12)
d.styles['Heading 1'].paragraph_format.space_before=Pt(8)
d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
d.styles['Heading 2'].font.size=Pt(10.5)
d.styles['Heading 2'].paragraph_format.space_before=Pt(7)
d.styles['Heading 2'].paragraph_format.space_after=Pt(4)
for style in d.styles:
 if style.type==1 and style.element.pPr is not None:
  for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)

def p(text='',style=None):return d.add_paragraph(text,style)
def heading(text):d.add_heading(text,level=1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
 t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for vals in rows:
  cells=t.add_row().cells
  for i,v in enumerate(vals):cells[i].text=v
 for n,row in enumerate(t.rows):
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   prop=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for tag,value in [('top',65),('bottom',65),('left',100),('right',100)]:
    e=OxmlElement('w:'+tag);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
   prop.append(margins)
   border=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
   prop.append(border)
   fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'163B5C' if n==0 else ('F0F5F8' if n%2==0 else 'FFFFFF'));prop.append(fill)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.08
    if i==0:para.alignment=1
    for run in para.runs:
     run.font.size=Pt(8.5 if len(headers)>3 and i==0 else 9 if len(headers)>3 else 9.5)
     run.font.color.rgb=RGBColor(255,255,255) if n==0 else RGBColor(0,0,0)
     run.bold=n==0
   keep=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(keep)
 spacer=p();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.add_run().font.size=Pt(2)
 return t

p('학교 자료 검색과 안내문 작성 수업계획안','Title')
p('부산보건고등학교 교사 AI 업무 연수','Subtitle')
p('학교 공개 자료를 NotebookLM에 등록해 필요한 내용을 검색하고, 교직원 안내문 작성에 활용하는 120분 과정입니다. RAG와 NotebookLM의 차이를 배우고 Word·PowerPoint·Excel 문서를 검토합니다.')
heading('연수 개요')
table(['항목','계획'],[
 ('대상','AI 도구로 자료를 생성해 본 교사 20명 내외를 기준으로 설계'),
 ('시간과 장소','총 120분 · 일시와 교육장은 학교 협의 후 확정'),
 ('강사','김세영 · ㈜파운드리머 기업부설연구소 연구소장'),
 ('준비물','노트북, Google 계정, 실습에 사용할 외부 AI 계정'),
 ('자료','학교규정집, 학교규정 개정 사항, 학업성적관리 지침 공개본')
],[1.2,5.9])
heading('학습 목표')
for text in ['RAG의 검색과 답변 생성 과정을 설명하고 NotebookLM과의 관계를 구분한다.','학교 공개 자료에서 필요한 데이터를 검색하고 인용 두 곳을 원문과 대조한다.','검색한 내용과 출처를 작성 요청에 포함해 문서 초안을 만들고 내용과 서식을 검토한다.']:
 p(text,'List Bullet')
heading('120분 진행 계획')
table(['시간','내용','학습 활동','결과물'],[
 ('00–10분','RAG와 NotebookLM','검색·생성 순서와 방식·제품 차이 설명','개념 비교'),
 ('10–25분','AI 도구 비교','계정별 기능·파일 작성 조건 확인, 실습 도구 선택','도구 선택'),
 ('25–55분','자료 등록과 검색','PDF 3종 등록 → 자료 선택 → 질문 → 인용 원문 대조','노트북·검색 기록'),
 ('55–65분','휴식','자료 처리와 계정 상태 점검',''),
 ('65–80분','스킬과 작성 요청','스킬 절차 참고, 검색 내용·출처·문서 형식 구성','작성 요청 1개'),
 ('80–110분','문서 작성 실습','Word·PPT·Excel 시연 후 한 형식 선택·작성·검토','검토 문서 1개'),
 ('110–120분','결과 확인','원문·날짜·수치·파일 형식 대조, 수정 사항 공유','검토 기록')
],[.85,1.28,3.57,1.4])

d.add_page_break()
heading('실습 진행안')
p('강사 시연 후 함께 따라 하고, 질문이나 문서 형식을 바꿔 혼자 실행한 뒤 결과를 확인합니다. 실제 학생 자료 대신 제공된 공개 자료와 가상 데이터를 사용합니다.')
for title,text in [
 ('학교 자료 등록과 검색','학교규정집·개정 사항·학업성적관리 지침을 구분해 등록합니다. 수행평가 결시 학생의 추가 평가 원칙을 질문하고 관련 자료를 선택했을 때와 해제했을 때의 결과를 비교합니다. 인용 두 곳을 열어 적용 대상·조건·예외와 기준 시점을 확인합니다.'),
 ('교직원 안내문 작성','확인한 발췌문, 자료명, 조항 또는 쪽, 적용 시점을 정리합니다. 사이트의 작성 양식에서 업무 목적과 필수 항목을 묶어 외부 AI에 전달합니다. 원칙·처리 절차·참고 자료·확인 필요 사항으로 구성한 A4 1쪽 초안을 작성합니다.'),
 ('Word PowerPoint Excel 검토','Word 또는 한글은 필수 항목·표·쪽 나눔을 확인합니다. PowerPoint는 같은 내용을 5장으로 구성하고 글자·표 넘침과 발표자 메모를 확인합니다. Excel은 가상 예산 CSV의 수량·단가·합계를 수식으로 작성하고 두 행을 직접 계산해 대조합니다. 참가자는 한 형식을 선택해 검토본을 저장합니다.')
]:
 d.add_heading(title,level=2);p(text)
heading('실습 질문과 작성 요청')
p('검색 질문  학교규정집에서 수행평가 결시 학생의 추가 평가 원칙을 3개 항목으로 정리해 줘. 각 항목에 인용을 붙이고 자료에 없는 내용은 확인 필요로 표시해 줘.')
p('작성 요청  확인한 발췌문과 자료명 조항 또는 쪽 적용 시점만 사용해 교직원 안내문 초안을 작성해 줘. 원칙·처리 절차·참고 자료·확인 필요 사항으로 구성하고 A4 1쪽 분량으로 정리해 줘. 제공되지 않은 담당자·일정·수치는 만들지 말아 줘.')
heading('완료 기준과 운영 준비')
for text in ['등록한 공개 자료의 제목과 적용 학년도를 확인한다.','인용 두 곳을 원문과 대조하고 날짜·숫자·적용 조건의 일치 여부를 기록한다.','검색 데이터를 포함한 작성 요청과 검토한 문서 파일 또는 복사 가능한 본문을 저장한다.','연수 전 계정 메뉴와 파일 업로드를 시험한다. 사용량이나 파일 생성이 제한되면 본문 슬라이드별 표 CSV를 문서 편집기에 옮겨 이어간다.','실제 학생 개인정보와 내부 업무 자료는 소속 기관이 허용하는 환경과 범위에서만 사용한다.']:
 p(text,'List Bullet')
heading('강의 사이트와 교재')
p(URL)
p('강의 교재 PDF, 공개 자료 묶음, 17단계 실습 화면, 도구 비교와 스킬 안내를 제공합니다.')

f=sec.footer.paragraphs[0];f.alignment=2
r=f.add_run('부산보건고 교사 AI 연수  |  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
d.core_properties.title='학교 자료 검색과 안내문 작성 수업계획안';d.core_properties.author='김세영'
path=OUT/'teacher-ai-lesson-plan.docx';d.save(path)
import shutil
shutil.copyfile(path,ROOT/'docs/downloads/teacher-ai-lesson-plan.docx')
print(path)
